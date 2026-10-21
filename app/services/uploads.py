"""Analyse de fichiers televerses (PDF, Word, texte, CSV).

Ephemere : le contenu reste en memoire, jamais sur disque ni en base.
Le fichier est une source NON FIABLE : validation du format, limites de
taille, filtre d'injection indirecte sur chaque extrait.
"""

import hashlib
import io
import json
import math
import os
import re
import time
import unicodedata
import uuid
import zipfile
from html.parser import HTMLParser
from collections import Counter
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Tuple

from app.services import vision
from rag.config import _dbg
from rag.generation import generate_answer
from rag.postcheck import verifier_fondement
from rag.security import MESSAGE_SORTIE_INVALIDE, detecter_extrait, sortie_fuite_prompt

try:
    import pymupdf
except ImportError:  # anciennes versions de PyMuPDF
    import fitz as pymupdf

MAX_OCTETS = 5 * 1024 * 1024
MAX_PAGES = 80
MAX_CARACTERES = 400_000
TAILLE_EXTRAIT = 1200
CHEVAUCHEMENT = 150
TOP_K = 6
TAILLE_PAGE_VIRTUELLE = 3000
IMAGES = (".png", ".jpg", ".jpeg", ".webp")
TEXTES = (".txt", ".md", ".log", ".tsv", ".csv", ".json", ".xml", ".html", ".htm")
EXTENSIONS = (".pdf", ".docx", ".xlsx", ".pptx") + TEXTES + IMAGES
DUREE_VIE_S = 30 * 60
MAX_PAR_UTILISATEUR = 3
MAX_PAGES_VISION = 6
MAX_CELLULES = 20000
RRF_K = 60
SEUIL_SEMANTIQUE = float(os.getenv("FICHIER_SEUIL_SEMANTIQUE", "0.30"))


class FichierRefuse(ValueError):
    """Fichier refuse : le message est affichable tel quel a l'utilisateur."""


@dataclass
class Fichier:
    id: str
    nom: str
    ext: str
    taille: int
    sha: str
    pages: int
    chunks: List[Dict[str, Any]]
    bloques: List[Dict[str, Any]] = field(default_factory=list)
    avertissements: List[str] = field(default_factory=list)
    pages_image: set = field(default_factory=set)
    emb: Any = None
    tenant_id: str = ""
    user_id: str = ""
    cree: float = field(default_factory=time.time)


# ----------------------------------------------------------------- lecture

def _nom_propre(nom: str) -> str:
    base = os.path.basename((nom or "").replace("\\", "/")) or "fichier"
    base = re.sub(r"[^\w .\-()]", "_", base, flags=re.UNICODE)
    return base[:100]


def _paginer(texte: str) -> List[Tuple[int, str]]:
    pages, courant, n = [], "", 1
    for ligne in texte.splitlines():
        if courant and len(courant) + len(ligne) > TAILLE_PAGE_VIRTUELLE:
            pages.append((n, courant))
            courant, n = "", n + 1
        courant += ligne + "\n"
    if courant.strip():
        pages.append((n, courant))
    return pages


def _lire_pdf(data: bytes):
    if not data.startswith(b"%PDF-"):
        raise FichierRefuse("Ce fichier n'est pas un vrai PDF.")
    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception:
        raise FichierRefuse("PDF illisible ou corrompu.")
    if doc.needs_pass or doc.is_encrypted:
        raise FichierRefuse("PDF protégé par mot de passe : retirez la protection puis réessayez.")
    if doc.page_count > MAX_PAGES:
        raise FichierRefuse("PDF trop long (%d pages, maximum %d)." % (doc.page_count, MAX_PAGES))
    pages, vides = [], []
    for i in range(doc.page_count):
        t = doc[i].get_text("text") or ""
        if len(t.strip()) < 20:
            vides.append(i)
        else:
            pages.append((i + 1, t))
    avert, pimg = [], set()
    if vides:  # pages scannees : lecture par le modele de vision (nombre limite)
        erreur = ""
        for i in vides[:MAX_PAGES_VISION]:
            try:
                jpeg = vision.preparer(doc[i].get_pixmap(dpi=150).tobytes("png"))
                texte = vision.transcrire(jpeg)
            except vision.ErreurVision as e:
                erreur = str(e)
                break
            if len(texte.strip()) >= 5:
                pages.append((i + 1, texte))
                pimg.add(i + 1)
        pages.sort()
        if pimg:
            avert.append("%d page(s) scannée(s) lue(s) par un modèle de vision : "
                         "vérifiez les chiffres importants sur l'original." % len(pimg))
        reste = len(vides) - len(pimg)
        if reste and pimg:
            avert.append("%d page(s) sans texte n'ont pas été lues (maximum %d pages scannées par fichier)."
                         % (reste, MAX_PAGES_VISION))
        if erreur and pimg:
            avert.append(erreur)
        if not pages:
            raise FichierRefuse(erreur or "Aucun texte lisible n'a été trouvé dans ce PDF.")
    return pages, avert, pimg


def _lire_image(data: bytes):
    ok = (data.startswith(b"\x89PNG\r\n\x1a\n") or data.startswith(b"\xff\xd8\xff")
          or (data[:4] == b"RIFF" and data[8:12] == b"WEBP"))
    if not ok:
        raise FichierRefuse("Ce fichier n'est pas une vraie image (PNG, JPG ou WebP).")
    try:
        texte = vision.transcrire(vision.preparer(data))
    except vision.ErreurVision as e:
        raise FichierRefuse(str(e))
    if len(texte.strip()) < 5:
        raise FichierRefuse("Aucun texte lisible n'a été trouvé dans cette image.")
    pages = _paginer(texte)
    return pages, ["Texte lu dans l'image par un modèle de vision : "
                   "vérifiez les chiffres importants sur l'original."], {p for p, _ in pages}


def _lire_docx(data: bytes) -> Tuple[List[Tuple[int, str]], List[str]]:
    if not data.startswith(b"PK"):
        raise FichierRefuse("Ce fichier n'est pas un vrai document Word (.docx).")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            infos = z.infolist()
            if len(infos) > 2000 or sum(i.file_size for i in infos) > 100 * 1024 * 1024:
                raise FichierRefuse("Document Word refusé (contenu décompressé trop volumineux).")
            if "word/document.xml" not in z.namelist():
                raise FichierRefuse("Ce fichier n'est pas un vrai document Word (.docx).")
        from docx import Document
        d = Document(io.BytesIO(data))
        lignes = [p.text for p in d.paragraphs if p.text.strip()]
        for tableau in d.tables:
            for ligne in tableau.rows:
                cellules = [c.text.strip() for c in ligne.cells]
                if any(cellules):
                    lignes.append(" | ".join(cellules))
    except FichierRefuse:
        raise
    except Exception:
        raise FichierRefuse("Document Word illisible ou corrompu.")
    texte = "\n\n".join(lignes)
    if not texte.strip():
        raise FichierRefuse("Ce document Word ne contient pas de texte.")
    return _paginer(texte), []


def _verifier_zip(data: bytes, marqueur: str, nom: str) -> None:
    if not data.startswith(b"PK"):
        raise FichierRefuse("Ce fichier n'est pas un vrai document %s." % nom)
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            infos = z.infolist()
            if len(infos) > 5000 or sum(i.file_size for i in infos) > 200 * 1024 * 1024:
                raise FichierRefuse("Document %s refusé (contenu décompressé trop volumineux)." % nom)
            if marqueur not in z.namelist():
                raise FichierRefuse("Ce fichier n'est pas un vrai document %s." % nom)
    except FichierRefuse:
        raise
    except zipfile.BadZipFile:
        raise FichierRefuse("Document %s illisible ou corrompu." % nom)


def _lire_xlsx(data: bytes):
    _verifier_zip(data, "xl/workbook.xml", "Excel (.xlsx)")
    lignes, cellules, tronque, feuilles = [], 0, False, 0
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for ws in wb.worksheets:
            feuilles += 1
            lignes.append("Feuille : %s" % ws.title)
            for row in ws.iter_rows(values_only=True):
                vals = ["" if v is None else str(v) for v in row]
                while vals and not vals[-1]:
                    vals.pop()
                if not vals:
                    continue
                cellules += len(vals)
                if cellules > MAX_CELLULES:
                    tronque = True
                    break
                lignes.append(" | ".join(vals))
            if tronque:
                break
    except FichierRefuse:
        raise
    except Exception:
        raise FichierRefuse("Classeur Excel illisible ou corrompu.")
    if cellules == 0:
        raise FichierRefuse("Ce classeur Excel ne contient aucune donnée.")
    avert = []
    if tronque:
        avert.append("Classeur tronqué : seules les %d premières cellules ont été lues." % MAX_CELLULES)
    return _paginer("\n".join(lignes)), avert


def _lire_pptx(data: bytes):
    _verifier_zip(data, "ppt/presentation.xml", "PowerPoint (.pptx)")
    pages = []
    try:
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        if len(prs.slides) > MAX_PAGES:
            raise FichierRefuse("Présentation trop longue (%d diapositives, maximum %d)." % (len(prs.slides), MAX_PAGES))
        for i, slide in enumerate(prs.slides, 1):
            morceaux = []
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    morceaux.append(shape.text_frame.text.strip())
                if getattr(shape, "has_table", False) and shape.has_table:
                    for row in shape.table.rows:
                        morceaux.append(" | ".join(c.text.strip() for c in row.cells))
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                note = slide.notes_slide.notes_text_frame.text.strip()
                if note:
                    morceaux.append("Notes : " + note)
            if morceaux:
                pages.append((i, "\n\n".join(morceaux)))
    except FichierRefuse:
        raise
    except Exception:
        raise FichierRefuse("Présentation PowerPoint illisible ou corrompue.")
    if not pages:
        raise FichierRefuse("Cette présentation ne contient pas de texte (images seulement).")
    return pages, []


def _lire_json(data: bytes):
    pages, avert = _lire_texte(data)
    texte = "".join(t for _, t in pages)
    try:
        texte = json.dumps(json.loads(texte), ensure_ascii=False, indent=1)
    except ValueError:
        return pages, avert
    return _paginer(texte), avert


_BLOCS = {"p", "div", "br", "li", "tr", "td", "th", "h1", "h2", "h3", "h4", "h5", "h6",
          "section", "article", "table", "ul", "ol"}


class _ExtracteurTexte(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.morceaux = []
        self._ignore = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self._ignore += 1
        elif tag in _BLOCS:
            self.morceaux.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style"):
            self._ignore = max(0, self._ignore - 1)
        elif tag in _BLOCS:
            self.morceaux.append("\n")

    def handle_data(self, data):
        if not self._ignore:
            self.morceaux.append(data)


def _lire_html(data: bytes):
    pages0, _ = _lire_texte(data)
    p = _ExtracteurTexte()
    try:
        p.feed("".join(t for _, t in pages0))
        p.close()
    except Exception:
        raise FichierRefuse("Fichier HTML/XML illisible.")
    lignes = [re.sub(r"[ \t]+", " ", l).strip() for l in "".join(p.morceaux).splitlines()]
    texte = "\n".join(l for l in lignes if l)
    if not texte:
        raise FichierRefuse("Aucun texte trouvé dans ce fichier.")
    return _paginer(texte), []


def _embeddings(chunks, encodeur):
    if encodeur is None or len(chunks) < 2:
        return None
    try:
        return encodeur.encode([c["content"] for c in chunks], batch_size=32,
                               normalize_embeddings=True, show_progress_bar=False)
    except Exception:
        return None


def _lire_texte(data: bytes) -> Tuple[List[Tuple[int, str]], List[str]]:
    if b"\x00" in data[:4096]:
        raise FichierRefuse("Ce fichier n'est pas un fichier texte.")
    for enc in ("utf-8-sig", "cp1252"):
        try:
            texte = data.decode(enc)
            break
        except UnicodeDecodeError:
            texte = ""
    if not texte.strip():
        raise FichierRefuse("Fichier vide ou encodage non reconnu.")
    return _paginer(texte), []


# --------------------------------------------------------------- decoupage

def _termes(s: str) -> List[str]:
    s = unicodedata.normalize("NFD", (s or "").lower())
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    out = []
    for t in re.findall(r"[a-z0-9]+", s):
        if t in _STOP or len(t) < 2:
            continue
        if len(t) > 3 and t[-1] in "sx":
            t = t[:-1]
        out.append(t)
    return out


_STOP = set((
    "le la les un une des du de d l et ou a au aux en dans sur pour par avec sans que qui quoi "
    "quel quelle quels quelles est sont ete etre il elle ils elles on ce cet cette ces se sa son "
    "ses leur leurs ne pas plus comment combien quand pourquoi dit selon document fichier pdf "
    "texte donne moi dis y a t ai as").split())


def _mk(texte: str, page: int, entete: str) -> Dict[str, Any]:
    contenu = texte
    if entete and not texte.startswith(entete):
        contenu = entete + "\n" + texte
    return {"content": contenu, "page": page, "tf": Counter(_termes(contenu))}


def _decouper(pages: List[Tuple[int, str]], entete: str = "") -> List[Dict[str, Any]]:
    chunks: List[Dict[str, Any]] = []
    sep = "\n" if entete else "\n\n"
    motif = r"\n" if entete else r"\n\s*\n"
    for num, texte in pages:
        courant = ""
        for u in [x.strip() for x in re.split(motif, texte) if x.strip()]:
            while len(u) > TAILLE_EXTRAIT:
                if courant:
                    chunks.append(_mk(courant, num, entete))
                    courant = ""
                chunks.append(_mk(u[:TAILLE_EXTRAIT], num, entete))
                u = u[TAILLE_EXTRAIT - CHEVAUCHEMENT:]
            if courant and len(courant) + len(u) + 2 > TAILLE_EXTRAIT:
                chunks.append(_mk(courant, num, entete))
                courant = ""
            courant = courant + sep + u if courant else u
        if courant:
            chunks.append(_mk(courant, num, entete))
    return chunks


def _ecarter_lignes(pages):
    """Retire les lignes suspectes (injection) en gardant le reste de la page."""
    propres, bloques = [], []
    for num, texte in pages:
        lignes = texte.split("\n")
        garde, retire = [], False
        for l in lignes:
            v = detecter_extrait(l) if l.strip() else None
            if v is not None and v.bloque:
                bloques.append({"page": num, "categorie": v.categorie, "regle": v.regle})
                retire = True
            else:
                garde.append(l)
        t = "\n".join(garde) if retire else texte
        if t.strip():
            propres.append((num, t))
    return propres, bloques


def charger(nom: str, data: bytes, encodeur: Any = None) -> Fichier:
    nom = _nom_propre(nom)
    ext = os.path.splitext(nom)[1].lower()
    if ext not in EXTENSIONS:
        raise FichierRefuse(
            "Format non pris en charge (%s). Formats acceptés : PDF, Word, Excel, PowerPoint, texte (TXT, MD, LOG, CSV, TSV, JSON, HTML, XML) et images (PNG, JPG, WebP)."
            % (ext or "sans extension"))
    if not data:
        raise FichierRefuse("Fichier vide.")
    if len(data) > MAX_OCTETS:
        raise FichierRefuse("Fichier trop volumineux (maximum %d Mo)." % (MAX_OCTETS // (1024 * 1024)))

    pimg: set = set()
    if ext == ".pdf":
        pages, avert, pimg = _lire_pdf(data)
    elif ext in IMAGES:
        pages, avert, pimg = _lire_image(data)
    elif ext == ".docx":
        pages, avert = _lire_docx(data)
    elif ext == ".xlsx":
        pages, avert = _lire_xlsx(data)
    elif ext == ".pptx":
        pages, avert = _lire_pptx(data)
    elif ext == ".json":
        pages, avert = _lire_json(data)
    elif ext in (".html", ".htm", ".xml"):
        pages, avert = _lire_html(data)
    else:
        pages, avert = _lire_texte(data)
    if sum(len(t) for _, t in pages) > MAX_CARACTERES:
        raise FichierRefuse("Contenu trop volumineux pour une analyse fiable.")

    entete = ""
    if ext in (".csv", ".tsv"):
        entete = next((l.strip() for l in pages[0][1].splitlines() if l.strip()), "")[:400]

    pages, bloques = _ecarter_lignes(pages)
    brut = _decouper(pages, entete)
    chunks = []
    for c in brut:  # le fichier est non fiable : meme filtre que les documents indexes
        v = detecter_extrait(c["content"])
        if v.bloque:
            bloques.append({"page": c["page"], "categorie": v.categorie, "regle": v.regle})
        else:
            chunks.append(c)
    if not chunks:
        raise FichierRefuse("Le contenu de ce fichier a été bloqué par la politique de sécurité.")
    if bloques:
        avert.append("%d passage(s) suspect(s) ont été écartés par la politique de sécurité." % len(bloques))
    if ext in (".csv", ".tsv", ".xlsx"):
        avert.append("Fichier tabulaire : l'assistant lit et cite les lignes, il ne fait pas de calculs.")

    emb = _embeddings(chunks, encodeur)
    return Fichier(
        id=uuid.uuid4().hex, nom=nom, ext=ext, taille=len(data),
        sha=hashlib.sha256(data).hexdigest(), pages=len(pages), chunks=chunks,
        bloques=bloques, avertissements=avert, pages_image=pimg, emb=emb)


# --------------------------------------------------------------- selection

_SYNTHESE = re.compile(
    r"resum|synthes|de quoi (parle|traite)|sujet (du|de ce)|points? cles?|principaux points|contenu (du|de ce)")


def _fold(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", (s or "").lower())
                   if unicodedata.category(c) != "Mn")


def selectionner(chunks: List[Dict[str, Any]], question: str, k: int = TOP_K,
                 emb: Any = None, qvec: Any = None) -> List[Dict[str, Any]]:
    """Mots (poids par rarete) + proximite de sens (si embeddings), fusionnes par rang."""
    n = len(chunks)
    if not n:
        return []
    q = _termes(question)
    if _SYNTHESE.search(_fold(question)) or not q:
        if n <= k:
            return list(chunks)
        return [chunks[round(i * (n - 1) / (k - 1))] for i in range(k)]
    df = Counter()
    for c in chunks:
        for t in set(c["tf"]):
            df[t] += 1
    lex = {}
    for i, c in enumerate(chunks):
        s = 0.0
        for t in set(q):
            f = c["tf"].get(t, 0)
            if f:
                s += math.log(1 + n / (1 + df[t])) * (1 + math.log(f))
        if s > 0:
            lex[i] = s
    sem = {}
    if emb is not None and qvec is not None:
        sims = emb @ qvec
        sem = {i: float(sims[i]) for i in range(n)}
    meilleure = max(sem.values()) if sem else -1.0
    if not lex and meilleure < SEUIL_SEMANTIQUE:
        _dbg("FICHIER selection : aucun passage (lex=0, sim_max=%.2f, seuil=%.2f)" % (meilleure, SEUIL_SEMANTIQUE))
        return []
    scores: Dict[int, float] = {}
    for r, i in enumerate(sorted(lex, key=lex.get, reverse=True)):
        scores[i] = scores.get(i, 0.0) + 1.0 / (RRF_K + r)
    for r, i in enumerate(i for i in sorted(sem, key=sem.get, reverse=True) if sem[i] >= SEUIL_SEMANTIQUE):
        scores[i] = scores.get(i, 0.0) + 1.0 / (RRF_K + r)
    top = sorted(scores, key=scores.get, reverse=True)[:k]
    _dbg("FICHIER selection : lex=%d, sim_max=%.2f, retenus=%d" % (len(lex), meilleure, len(top)))
    return [chunks[i] for i in sorted(top)]


# ----------------------------------------------------------------- reponse

def _score(fo: Any) -> Optional[float]:
    if isinstance(fo, dict):
        fo = fo.get("score")
    return float(fo) if isinstance(fo, (int, float)) else None


_RE_ABSENCE = re.compile(
    r"(aucun(?:e)?\s+(?:des\s+|de\s+ces\s+)?(?:extraits?|passages?|informations?|donn[ée]es?)"
    r"|ne\s+(?:mentionn|contien|figur|pr[ée]cis|donn|trouv|parl)\w*"
    r"|n['’]est\s+(?:pas\s+)?(?:mentionn|indiqu|pr[ée]cis|disponible)\w*"
    r"|pas\s+d['’](?:information|donn[ée]e)"
    r"|information\s+(?:absente|introuvable)"
    r"|je\s+ne\s+trouve\s+pas)", re.I)
_RE_RENVOI = re.compile(r"\[\s*(?:Source\s*)?\d+")


def _est_un_refus(rep: str) -> bool:
    """Reponse courte, sans aucun renvoi [n], qui dit que l'information est absente.
    Une vraie reponse porte toujours des renvois : celle-ci ne doit ni citer de
    source ni recevoir de score de fondement."""
    return len(rep) < 400 and not _RE_RENVOI.search(rep) and bool(_RE_ABSENCE.search(rep))


def repondre(f: Fichier, question: str, role: str, encodeur: Any = None) -> Dict[str, Any]:
    qvec = None
    if f.emb is not None and encodeur is not None:
        try:
            qvec = encodeur.encode([question], normalize_embeddings=True, show_progress_bar=False)[0]
        except Exception:
            qvec = None
    sel = selectionner(f.chunks, question, emb=f.emb if qvec is not None else None, qvec=qvec)
    if not sel:
        return {"answer": "Je ne trouve pas cette information dans le fichier « %s »." % f.nom,
                "fondement": None, "sources": [], "trouve": False}
    rag = [{
        "content": c["content"], "fichier_source": f.nom, "page": c["page"],
        "entreprise": "Fichier : " + f.nom, "type_document": "fichier_utilisateur",
        "type": "fichier_utilisateur", "departement": "", "score": 1.0,
        "demo": False, "verifie": False,
    } for c in sel]
    rep = generate_answer(question, rag, role)
    if isinstance(rep, dict):
        rep = rep.get("answer", "")
    rep = str(rep or "")
    if sortie_fuite_prompt(rep):
        return {"answer": MESSAGE_SORTIE_INVALIDE, "fondement": None, "sources": [], "trouve": True}
    if _est_un_refus(rep):
        return {"answer": "Je ne trouve pas cette information dans le fichier « %s »." % f.nom,
                "fondement": None, "sources": [], "trouve": False}
    cites = sorted({int(m) for m in re.findall(r"\[(?:Source\s*)?(\d+)", rep) if 0 < int(m) <= len(rag)})
    cites = cites or list(range(1, len(rag) + 1))
    sources = [{"n": i, "fichier": f.nom, "page": rag[i - 1]["page"], "entreprise": "Fichier téléversé",
                "type": "temporaire", "demo": False, "verifie": False} for i in cites]
    fond = _score(verifier_fondement(rep, rag))
    if f.pages_image and any(c["page"] in f.pages_image for c in sel):
        fond = None  # texte lu par un modele de vision : le score ne prouverait rien
    return {"answer": rep, "fondement": fond, "sources": sources, "trouve": True}
