import os
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

FILL = " Cette disposition s'applique à tous les collaborateurs de la société Atlas Demo, quel que soit leur site de rattachement."
SUJETS = [
 ("Congés", "Les congés annuels sont de 22 jours ouvrables par an. Toute demande doit être déposée au moins deux semaines à l'avance auprès du responsable. Le solde non pris au 31 décembre est reporté jusqu'au 31 mars de l'année suivante."),
 ("Rémunération", "La grille salariale comprend cinq niveaux. Le salaire d'un technicien débutant est de 6 500 MAD brut par mois. Une revue annuelle des salaires a lieu en janvier."),
 ("Télétravail", "Le télétravail est autorisé deux jours par semaine après six mois d'ancienneté. Le collaborateur doit disposer d'une connexion internet stable et d'un espace de travail adapté."),
 ("Déplacements", "Les frais de déplacement sont remboursés sur justificatifs dans un délai de 15 jours. Le plafond journalier pour un repas est de 150 MAD."),
 ("Sécurité informatique", "Les mots de passe doivent contenir au moins 12 caractères et être renouvelés tous les 90 jours. L'accès aux serveurs nécessite une authentification à deux facteurs."),
 ("Formation", "Le budget de formation 2026 est de 250 000 MAD. Chaque collaborateur peut suivre au maximum cinq jours de formation par an."),
 ("Recrutement", "Le processus de recrutement comprend trois entretiens. La période d'essai est de trois mois renouvelable une fois pour les cadres."),
 ("Évaluation", "L'évaluation annuelle des performances a lieu en novembre. Elle donne lieu à un objectif chiffré pour l'année suivante."),
 ("Avantages", "La mutuelle couvre 80 % des frais médicaux. Les tickets restaurant sont de 40 MAD par jour travaillé."),
 ("Départ", "Le préavis est de trois mois pour les cadres et d'un mois pour les employés."),
]
os.makedirs("exemples", exist_ok=True)
d = Document(); d.add_heading("Guide RH - Atlas Demo (fictif)", 1)
for titre, texte in SUJETS:
    d.add_heading(titre, 2); d.add_paragraph(texte + FILL * 3)
d.save("exemples/guide_rh_long.docx")

wb = Workbook(); ws = wb.active; ws.title = "Ventes"
ws.append(["region", "trimestre", "ventes_mad"])
for r in (("Casablanca", "T1", 120000), ("Rabat", "T1", 80000), ("Tanger", "T2", 95000)):
    ws.append(list(r))
wb.save("exemples/ventes.xlsx")

prs = Presentation()
s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = "Plan 2027"
s.placeholders[1].text = "Ouverture de deux agences à Agadir et Oujda. Budget d'investissement : 4 millions MAD."
prs.save("exemples/plan_2027.pptx")
print("OK : guide_rh_long.docx, ventes.xlsx, plan_2027.pptx")
